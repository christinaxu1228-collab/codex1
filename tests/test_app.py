import io
import asyncio
import struct
from pathlib import Path
from openpyxl import Workbook, load_workbook
import app.main as main
from app.main import validate, export_workbook, prepare_products, TemplateError, verify_export, index

def test_validation_duplicate_and_required():
    rows=[{"category":"女士短连衣裙","title":"A","description_en":"dress","sku":"X","color":"Red","size":"M","price":"100","stock":"2","image_url":"http://x"},{"sku":"X"}]
    errors=validate(rows)
    assert any(e["message"] == "重复 SKU" for e in errors)
    assert any(e["field"] == "title" for e in errors)

def test_category_code_is_not_price_range():
    row = {"category":"女士短连衣裙","title":"A","description_en":"dress","sku":"PRICE-1","color":"Red","size":"M","price":"2000","stock":"2","image_url":"http://x"}
    assert validate([row]) == []
    row["price"] = "not-a-number"
    assert any(e["field"] == "price" and "数字" in e["message"] for e in validate([row]))

def test_default_price_and_manual_price_are_preserved():
    rows = prepare_products([{"category":"女士连体裤", "sku":"A", "price":""}, {"category":"女士两件套", "sku":"B", "price":"12.345"}])
    assert rows[0]["price"] == "200.00"
    assert rows[1]["price"] == "12.35"

def test_custom_default_price():
    assert prepare_products([{"price":""}], "188.8")[0]["price"] == "188.80"

def test_three_categories_accept_default_price_and_export():
    rows = []
    for category, sku in [("女士两件套", "SET-1"), ("女士短连衣裙", "DR-1"), ("女士连体裤", "JP-1")]:
        rows.append(prepare_products([{"category": category, "title": "A", "description_en": "item", "sku": sku, "color": "Red", "size": "M", "price": "", "stock": "1", "image_url": "https://img/x"}])[0])
    assert [r["price"] for r in rows] == ["200.00", "200.00", "200.00"]
    assert validate(rows) == []

def template_bytes():
    wb=Workbook(); wb.remove(wb.active)
    for index,name in enumerate(["女士短连衣裙 67-1727", "女士两件套 95-1780", "女士连体裤 154-1860"]):
        ws=wb.create_sheet(name)
        for r in range(1,7): ws.cell(r,1,f"保留行{r}" if index == 0 else None)
        for i,h in enumerate(["category","商品名称","英文描述","supplier_sku","颜色","尺码","cost","仓库库存","商品图片"],1): ws.cell(2,i,h)
        for i,h in enumerate(["category","default_product_name[en]","default_product_desc[en]","supplier_sku","color","size","cost","warehouse_stock","main_image"],1): ws.cell(3,i,h)
    b=io.BytesIO(); wb.save(b); return b.getvalue()

def test_export_preserves_header_rows():
    data=export_workbook(template_bytes(), [{"category":"女士短连衣裙","title":"A","description_en":"dress","sku":"X","color":"Red","size":"M","price":"100","stock":"2","image_url":"http://x"}])
    ws=load_workbook(io.BytesIO(data)).active
    assert ws["A1"].value == "保留行1" and ws["B7"].value == "A" and ws["G7"].value == 100

def full_template_bytes():
    wb=Workbook(); wb.remove(wb.active)
    for name in ["女士两件套 95-1780", "女士短连衣裙 67-1727", "女士连体裤 154-1860", "辅助数据"]:
        ws=wb.create_sheet(name)
        for i,h in enumerate(["supplier_sku","default_product_name[en]","default_product_desc[en]","cost","color","size","warehouse_stock","main_image"],1): ws.cell(2,i,h); ws.cell(3,i,h)
    b=io.BytesIO(); wb.save(b); return b.getvalue()

def test_three_category_sheets_and_cost_price():
    rows=[]
    for category, sku in [("女士两件套","S1"),("女士短连衣裙","S2"),("女士连体裤","S3")]:
        rows.append({"category":category,"title":"A","description_en":"D","sku":sku,"price":"200.00","color":"Red","size":"M","stock":"1","image_url":"https://x"})
    data=export_workbook(full_template_bytes(), rows)
    wb=load_workbook(io.BytesIO(data))
    for name in ["女士两件套 95-1780","女士短连衣裙 67-1727","女士连体裤 154-1860"]:
        assert wb[name]["D7"].value == 200
        assert wb[name]["A7"].value in {"S1","S2","S3"}
        assert wb[name]["A2"].value == "supplier_sku"

def test_export_readback_verification():
    rows=[{"category":"女士短连衣裙","title":"A","description_en":"D","sku":"S1","price":"200.00","color":"Red","size":"M","stock":"1","image_url":"https://x"}]
    template=template_bytes(); result=export_workbook(template, rows)
    assert verify_export(template, result, rows) == []

def test_missing_category_sheet_blocks_export():
    wb=Workbook(); wb.active.title="女士短连衣裙 67-1727"; b=io.BytesIO(); wb.save(b)
    try:
        export_workbook(b.getvalue(), [])
    except TemplateError:
        pass

def test_inspect_endpoint():
    html = index()
    assert "SHEIN" in html and "默认申报供货价" in html

def test_draft_save_and_reload_preserves_edits(tmp_path):
    original_path = main.DB_PATH
    try:
        main.DB_PATH = Path(tmp_path) / "drafts.sqlite3"
        main.init_db()
        rows = [{"sku": "SKU-1", "supplier_code": "SPU-1", "price": "200.00", "pricing_status": "需要调整报价", "pricing_note": "保留备注", "attr_160": "棉"}]
        result = asyncio.run(main.save_drafts({"products": rows}))
        assert result == {"saved": 1}
        restored = main.load_drafts()["products"]
        assert restored == rows
    finally:
        main.DB_PATH = original_path

def test_main_attribute_value_survives_prepare_and_export():
    row = prepare_products([{"category":"女士连体裤", "supplier_code":"SPU-1", "sku":"TEST-JUMP-001", "title":"Jumpsuit", "description_en":"D", "price":"200", "main_attribute":"颜色", "main_attribute_value":"Black", "attr_39":"中弹", "attr_160":"棉"}])[0]
    assert row["main_attribute_value"] == "Black"

def test_final_price_wins_over_legacy_declared_price():
    rows = [{"category":"女士短连衣裙", "sku":"PRICE-1", "title":"A", "description_en":"D", "price":"200.00", "declared_price":"22.99", "color":"黑色", "size":"M", "stock":"1", "image_url":"https://x"}]
    result = export_workbook(full_template_bytes(), rows)
    wb = load_workbook(io.BytesIO(result))
    assert wb["女士短连衣裙 67-1727"]["D7"].value == 200

def test_different_final_prices_round_trip_all_categories():
    prices = {"女士两件套": "200.00", "女士短连衣裙": "18.75", "女士连体裤": "31.20"}
    rows = [{"category": c, "sku": "P-" + c, "title":"A", "description_en":"D", "price":p, "declared_price":"19.99", "color":"黑色", "size":"M", "stock":"1", "image_url":"https://x"} for c,p in prices.items()]
    template = full_template_bytes(); result = export_workbook(template, rows)
    assert verify_export(template, result, rows) == []

def test_fabric_elasticity_values_are_restricted():
    base={"category":"女士连体裤","title":"A","description_en":"D","sku":"E1","price":"200","attr_39":"低弹","color":"黑色","size":"M","stock":"1","image_url":"https://x"}
    assert validate([base]) == []
    base["attr_39"]="未知弹性"
    assert any(e["field"]=="attr_39" for e in validate([base]))

def test_fabric_elasticity_is_preserved_in_draft_payload():
    row=prepare_products([{"sku":"E2","price":"200","attr_39":"高弹"}])[0]
    assert row["attr_39"] == "高弹"

def test_official_template_exposes_image_fields():
    template = Path("/workspace/attachments/cec88a87-2a53-4b1f-87fe-a463b14793c8/批量添加商品模板 (2).xlsx")
    if not template.exists():
        return
    wb = load_workbook(template)
    metadata = main.template_metadata(wb)
    for category in main.CATEGORIES:
        fields = metadata[category]["image_fields"]
        assert fields[:3] == ["main_img", "square_img", "piece_img"]
        assert "detail_img_10" in fields and "sku_image" in fields

def test_local_image_asset_is_blocked_from_formal_export():
    row = {"category":"女士短连衣裙", "sku":"IMG-1", "title":"A", "description_en":"D", "price":"200",
           "color":"黑色", "size":"M", "stock":"1", "images":{"main_img": {"path":"product_images/x.png", "url":"", "width":900, "height":1200, "mime":"image/png", "size":100}}}
    wb = load_workbook(io.BytesIO(full_template_bytes()))
    for ws in [wb[x] for x in ["女士两件套 95-1780", "女士短连衣裙 67-1727", "女士连体裤 154-1860"]]:
        ws.cell(2, 9, "首图"); ws.cell(3, 9, "main_img")
    b=io.BytesIO(); wb.save(b)
    try:
        export_workbook(b.getvalue(), [row])
    except TemplateError as exc:
        assert "公开图片 URL" in str(exc)

def test_png_dimension_parser_and_image_rules():
    raw = b"\x89PNG\r\n\x1a\n" + b"\x00" * 8 + struct.pack(">II", 900, 1200) + b"\x00" * 20
    assert main.image_dimensions(raw, ".png")[:2] == (900, 1200)
    asset = {"url":"https://cdn.example/x.png", "width":900, "height":1200, "mime":"image/png", "size":100}
    assert main.validate_image_for_field(asset, "main_img") == []
    assert main.validate_image_for_field({**asset, "width":80, "height":80}, "main_img")

def image_template_with_slots():
    wb = load_workbook(io.BytesIO(full_template_bytes()))
    for ws in [wb[x] for x in ["女士两件套 95-1780", "女士短连衣裙 67-1727", "女士连体裤 154-1860"]]:
        for col, field in enumerate(["main_img", "detail_img_1", "detail_img_2"], 9):
            ws.cell(2, col, field); ws.cell(3, col, field)
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def test_display_images_require_two_and_map_in_user_order():
    base = {"category":"女士短连衣裙", "sku":"IMG-2", "title":"A", "description_en":"D", "price":"200", "color":"黑色", "size":"M", "stock":"1"}
    one = {**base, "images":{"display":["https://cdn.example/a.jpg"]}}
    try:
        export_workbook(image_template_with_slots(), [one])
        assert False, "one display image should be rejected"
    except TemplateError as exc:
        assert "少于 2 张" in str(exc)
    two = {**base, "images":{"display":["https://cdn.example/a.jpg", "https://cdn.example/b.jpg"]}}
    result = export_workbook(image_template_with_slots(), [two])
    ws = load_workbook(io.BytesIO(result))["女士短连衣裙 67-1727"]
    assert ws["I7"].value == "https://cdn.example/a.jpg"
    assert ws["J7"].value == "https://cdn.example/b.jpg"

def test_display_image_asset_requires_three_four_ratio_but_no_fixed_category():
    asset = {"url":"https://cdn.example/a.jpg", "width":1000, "height":1000, "mime":"image/jpeg", "size":100}
    errors = main.validate_image_for_field(asset, "main_img")
    assert any("3:4" in error for error in errors)

